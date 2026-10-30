import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Set page margins (1.8 cm / ~0.7 inches for compact multi-entry form)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Style: Default font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("DOCUMENTO DE EXENCIÓN DE RESPONSABILIDAD,")
    run_title.bold = True
    run_title.font.size = Pt(13.5)
    run_title.font.name = 'Calibri'
    run_title.font.color.rgb = RGBColor(0x8B, 0x00, 0x00) # Dark red / sangre

    p_title2 = doc.add_paragraph()
    p_title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title2.paragraph_format.space_before = Pt(0)
    p_title2.paragraph_format.space_after = Pt(4)
    run_title2 = p_title2.add_run("RENUNCIA VITAL Y CONSENTIMIENTO INFORMADO")
    run_title2.bold = True
    run_title2.font.size = Pt(13.5)
    run_title2.font.name = 'Calibri'
    run_title2.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

    # Subtitle / Docket number
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    run_sub = p_sub.add_run("EXPEDIENTE PARANORMAL DAW-666/2026 · AULA DE INFORMÁTICA · EDICIÓN HALLOWEEN")
    run_sub.font.size = Pt(8.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Intro
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(4)
    p_intro.paragraph_format.line_spacing = 1.15
    run_manif = p_intro.add_run("MANIFIESTO Y JURAMENTO DEL VISITANTE:\n")
    run_manif.bold = True
    run_manif.font.size = Pt(9.5)
    run_intro = p_intro.add_run(
        "El abajo firmante comparece de manera voluntaria, presuntamente en su sano juicio y libre de coacción a las puertas "
        "de este laboratorio. Con plena advertencia del peligro sensorial al que se expone, "
    )
    run_intro_bold = p_intro.add_run("DECLARA, ASUME Y ACEPTA BAJO JURAMENTO MORTAL ")
    run_intro_bold.bold = True
    p_intro.add_run("las siguientes cláusulas vinculantes:")

    # Clauses
    clauses = [
        ("1. Aptitud física y riesgo de colapso: ",
         "El firmante declara formalmente no padecer afecciones cardíacas, debilidad de esfínteres ni pánico patológico. "
         "Reconoce que los sucesos en este laboratorio pueden provocar taquicardias repentinas, gritos incontrolados, espasmos "
         "musculares, sudor frío y pesadillas recurrentes durante las próximas semanas."),
        
        ("2. Exoneración total al gremio de DAW: ",
         "Los estudiantes y docentes de Desarrollo de Aplicaciones Web quedan total, legal, moral y espiritualmente eximidos "
         "de responsabilidad civil o médica ante cualquier daño físico (caídas de la silla, hiperventilación, golpes al retroceder) "
         "o daño psicológico (crisis de fe en la informática o pérdida temporal del alma)."),
        
        ("3. Cesión del destino y fenómenos espectrales: ",
         "Ni el centro docente ni el equipo de programación responden de las entidades, espectros de memoria RAM o presencias malignas "
         "que se manifiesten a través del teclado o la pantalla. Cualquier anomalía espectral será catalogada como fallo ajeno al sistema."),
        
        ("4. Renuncia a indemnizaciones terrenales: ",
         "En caso de shock traumático, mutismo temporal o desaparición física del sujeto, los herederos legales renuncian expresamente "
         "a emprender acciones judiciales o eclesiásticas contra los organizadores."),
        
        ("5. Cláusula de no arrepentimiento: ",
         "«Entiendo que nadie me ha obligado a entrar. Si entro, es bajo mi exclusiva cuenta y riesgo. Si salgo temblando, yo me lo he buscado».")
    ]

    for title, text in clauses:
        p_cl = doc.add_paragraph()
        p_cl.paragraph_format.space_before = Pt(1)
        p_cl.paragraph_format.space_after = Pt(3)
        p_cl.paragraph_format.line_spacing = 1.1
        run_t = p_cl.add_run(title)
        run_t.bold = True
        run_t.font.size = Pt(8.5)
        run_t.font.color.rgb = RGBColor(0x33, 0x11, 0x11)
        run_txt = p_cl.add_run(text)
        run_txt.font.size = Pt(8.5)

    # Signature notice
    p_sig_note = doc.add_paragraph()
    p_sig_note.paragraph_format.space_before = Pt(6)
    p_sig_note.paragraph_format.space_after = Pt(4)
    run_sn = p_sig_note.add_run("REGISTRO DE VALIENTES (FIRMA OBLIGATORIA ANTES DE CRUZAR EL UMBRAL):")
    run_sn.bold = True
    run_sn.font.size = Pt(9)
    run_sn.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

    # Table
    table = doc.add_table(rows=13, cols=5)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Nº", "Nombre y Apellidos", "Curso / Grupo", "¿A quién avisar si colapsas?", "Firma del Declarante"]
    col_widths = [Inches(0.4), Inches(2.3), Inches(1.2), Inches(1.8), Inches(1.4)]

    # Header row formatting
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(8.5)
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(hdr_cells[i], "1F1A24")  # Dark gothic header
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Row data
    for row_idx in range(1, 13):
        row_cells = table.rows[row_idx].cells
        # Number
        row_cells[0].text = f"{row_idx:02d}"
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[0].paragraphs[0].runs[0].font.size = Pt(8)
        row_cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        
        # Set zebra striping or clean white
        bg_color = "F9F8FA" if row_idx % 2 == 0 else "FFFFFF"
        for i in range(5):
            set_cell_background(row_cells[i], bg_color)
            row_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Set row height
        trPr = table.rows[row_idx]._tr.get_or_add_trPr()
        trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="380" w:hRule="atLeast"/>')
        trPr.append(trHeight)

    # Set column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Footer note
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_footer.paragraph_format.space_before = Pt(4)
    run_fn = p_footer.add_run("* Archívese en los anales del Departamento de Informática. La dirección no devuelve almas.")
    run_fn.font.italic = True
    run_fn.font.size = Pt(7.5)
    run_fn.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    output_path = r"c:\xampp\htdocs\DAWlloween\assets\documentos\Exencion_Responsabilidad_DAWlloween.docx"
    doc.save(output_path)
    print(f"Document created successfully at: {output_path}")

if __name__ == "__main__":
    create_document()

